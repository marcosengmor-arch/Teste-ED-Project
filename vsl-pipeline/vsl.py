#!/usr/bin/env python3
"""Pipeline VSL: vídeo → transcrição → tradução → doc com prints (padrão SODA TIDE).

Fases mecânicas do PROMPT-Pipeline-VSL.md. As fases de julgamento (falantes,
blocos, tradução, legendas dos prints) são feitas pelo Claude e gravadas em
doc.json, que `docx` monta no layout do doc-modelo (ver MODELO.md).

Chave da AssemblyAI: network secret do ambiente cloud. O proxy da Anthropic
injeta o header `authorization` para api.assemblyai.com; a sessão nunca vê a
chave e o script envia as requisições sem credencial. (Fallback local:
variável ASSEMBLYAI_API_KEY.) Nunca em arquivo, log ou chat.

Subcomandos:
  check                                   testa chave + acesso à AssemblyAI
  prep   VIDEO --name N                   pasta VSL-Pipeline-N/, duração, resolução
  audio  VIDEO --name N                   audio_vsl.mp3 128k
  transcribe --name N [--lang en|pt]      AssemblyAI com speaker_labels
  frames VIDEO --name N [--every 60] [--extra ...]   prévia automática (troca de falante + cada N s)
  frames VIDEO --name N --doc doc.json    extrai os prints declarados no doc.json (nome bNN_desc_HH-MM-SS.jpg)
  sheet  --name N [--cols 4]              contact sheets dos frames para legendar olhando
  docx   --name N [--doc doc.json]        monta o .docx no layout-modelo
  verify --name N [--doc doc.json]        checagem final (contagens, ordem, arquivos)
"""
import argparse
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path

AAI = "https://api.assemblyai.com/v2"


# ---------- utilidades ----------
def work(name):
    return Path(f"VSL-Pipeline-{name}")


def ts_doc(sec):
    """Formato do doc-modelo: m:ss, minutos sem limite (80:57)."""
    sec = int(round(sec))
    return f"{sec // 60}:{sec % 60:02d}"


def ts_file(sec):
    """Formato do nome de arquivo: HH-MM-SS."""
    sec = int(round(sec))
    return f"{sec // 3600:02d}-{(sec % 3600) // 60:02d}-{sec % 60:02d}"


def ts_to_sec(ts):
    parts = [int(p) for p in str(ts).strip().replace("-", ":").split(":")]
    while len(parts) < 3:
        parts.insert(0, 0)
    h, m, s = parts
    return h * 3600 + m * 60 + s


def probe(video):
    out = subprocess.run(
        ["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries",
         "stream=width,height:format=duration", "-of", "json", video],
        capture_output=True, text=True, check=True).stdout
    j = json.loads(out)
    st = (j.get("streams") or [{}])[0]
    return float(j["format"]["duration"]), st.get("width"), st.get("height")


def headers():
    """Sem credencial por padrão: o network secret injeta o header no proxy."""
    key = os.environ.get("ASSEMBLYAI_API_KEY", "").strip()
    return {"authorization": key} if key else {}


def load_doc(a):
    w = work(a.name)
    p = Path(a.doc) if a.doc else w / "doc.json"
    d = json.loads(p.read_text(encoding="utf-8"))
    d.setdefault("frames_dir", str(w / "frames"))
    return d


# ---------- fases 1–3 ----------
def cmd_check(a):
    import requests
    try:
        r = requests.get(f"{AAI}/transcript", headers=headers(), params={"limit": 1}, timeout=20)
    except requests.RequestException as e:
        print(f"REDE: não alcançou api.assemblyai.com -> {e}")
        print("O secret não está ativo nesta sessão: confira 'Sites permitidos' = api.assemblyai.com "
              "no network secret, ou abra uma sessão nova.")
        return 1
    if r.status_code == 200:
        print("OK: chave válida e api.assemblyai.com acessível.")
        return 0
    if r.status_code in (401, 403):
        modo = "variável ASSEMBLYAI_API_KEY" if headers() else "network secret do ambiente"
        print(f"CHAVE: AssemblyAI respondeu {r.status_code} usando {modo}.")
        print("Confira o secret: header Authorization sem prefixo, valor = chave.")
        return 1
    print(f"Resposta inesperada {r.status_code}: {r.text[:200]}")
    return 1


def cmd_prep(a):
    if not Path(a.video).is_file():
        sys.exit(f"Vídeo não encontrado: {a.video}")
    w = work(a.name)
    (w / "frames").mkdir(parents=True, exist_ok=True)
    dur, wd, ht = probe(a.video)
    meta = {"video": str(Path(a.video).resolve()), "arquivo": Path(a.video).name,
            "duracao_s": dur, "duracao": ts_doc(dur), "largura": wd, "altura": ht}
    (w / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2))
    print(f"Pasta {w}/ criada. {meta['arquivo']}: {meta['duracao']}, {wd}x{ht}, "
          f"{os.path.getsize(a.video) / 1e6:.1f} MB")
    return 0


def cmd_audio(a):
    w = work(a.name)
    w.mkdir(parents=True, exist_ok=True)
    out = w / "audio_vsl.mp3"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", a.video, "-vn",
                    "-acodec", "libmp3lame", "-ab", "128k", "-ar", "44100", str(out)], check=True)
    print(f"Áudio: {out} ({out.stat().st_size / 1e6:.1f} MB)")
    return 0


def cmd_transcribe(a):
    import requests
    h = headers()
    w = work(a.name)
    audio = w / "audio_vsl.mp3"
    if not audio.is_file():
        sys.exit(f"{audio} não existe. Rode `audio` antes.")
    with open(audio, "rb") as f:
        up = requests.post(f"{AAI}/upload", headers=h, data=f, timeout=900)
    up.raise_for_status()
    j = requests.post(f"{AAI}/transcript", headers=h, timeout=60, json={
        "audio_url": up.json()["upload_url"], "language_code": a.lang,
        "speaker_labels": True, "punctuate": True, "format_text": True})
    j.raise_for_status()
    tid = j.json()["id"]
    print(f"Job {tid} enviado. Aguardando...")
    while True:
        r = requests.get(f"{AAI}/transcript/{tid}", headers=h, timeout=60).json()
        if r["status"] == "completed":
            break
        if r["status"] == "error":
            sys.exit(f"Erro AssemblyAI: {r.get('error')}")
        time.sleep(15)
    (w / "transcricao_raw.json").write_text(json.dumps(r, ensure_ascii=False, indent=2))
    utts = [{"speaker": u["speaker"], "start": ts_doc(u["start"] / 1000), "end": ts_doc(u["end"] / 1000),
             "start_ms": u["start"], "end_ms": u["end"], "text": u["text"]}
            for u in (r.get("utterances") or [])]
    (w / "segments.json").write_text(json.dumps(utts, ensure_ascii=False, indent=2))
    with open(w / "transcricao.md", "w", encoding="utf-8") as f:
        f.write(f"# Transcrição bruta — {a.name}\n\n")
        for u in utts:
            f.write(f"**[{u['start']} → {u['end']}] Falante {u['speaker']}:** {u['text']}\n\n")
    print(f"Concluído: {len(utts)} falas, falantes {sorted({u['speaker'] for u in utts})}, "
          f"modelo {r.get('speech_model')}, {r.get('audio_duration')}s. "
          f"Arquivos: transcricao_raw.json, segments.json, transcricao.md")
    return 0


# ---------- fase 6: frames ----------
def extract_frame(video, sec, out):
    for delta in (0, -2, 2):
        t = sec + delta
        if t < 0:
            continue
        r = subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-ss", str(t), "-i", video,
                            "-frames:v", "1", "-q:v", "2", str(out)], capture_output=True)
        if r.returncode == 0 and out.is_file() and out.stat().st_size > 0:
            return t
    return None


def auto_timestamps(segments, every, extra, total_s):
    pts, prev, last = {0}, None, 0
    for u in segments:
        s = u["start_ms"] // 1000
        if u["speaker"] != prev:
            pts.add(s)
            last, prev = s, u["speaker"]
        e = u["end_ms"] // 1000
        while e - last > every:
            last += every
            pts.add(last)
    pts.update(ts_to_sec(t) for t in extra)
    return sorted(p for p in pts if p < total_s)


def cmd_frames(a):
    w = work(a.name)
    fdir = w / "frames"
    fdir.mkdir(parents=True, exist_ok=True)
    total, _, _ = probe(a.video)
    missing, done = [], 0
    if a.doc:
        d = load_doc(a)
        fdir = Path(d["frames_dir"])
        fdir.mkdir(parents=True, exist_ok=True)
        for b in d["blocos"]:
            for mb in b["mini"]:
                for pr in mb.get("prints", []):
                    sec = ts_to_sec(pr["ts"])
                    if sec >= total:
                        missing.append(pr["arquivo"]); continue
                    if extract_frame(a.video, sec, fdir / pr["arquivo"]) is None:
                        missing.append(pr["arquivo"])
                    else:
                        done += 1
    else:
        seg_file = w / "segments.json"
        segments = json.loads(seg_file.read_text()) if seg_file.is_file() else []
        if not segments and not a.extra:
            sys.exit("Sem segments.json e sem --extra: nada para extrair.")
        index = {}
        for sec in auto_timestamps(segments, a.every, a.extra, total):
            out = fdir / f"auto_{ts_file(sec)}.jpg"
            got = extract_frame(a.video, sec, out)
            index[ts_doc(sec)] = None if got is None else str(out)
            if got is None:
                missing.append(ts_doc(sec))
            else:
                done += 1
        (fdir / "auto_index.json").write_text(json.dumps(index, ensure_ascii=False, indent=2))
    print(f"Frames: {done} extraídos em {fdir}/, {len(missing)} indisponíveis {missing}")
    return 0


def cmd_sheet(a):
    from PIL import Image, ImageDraw
    fdir = work(a.name) / "frames"
    files = sorted(p for p in fdir.glob("*.jpg") if not p.name.startswith("contact_"))
    if not files:
        sys.exit("Nenhum frame em frames/.")
    cols, cell_w, per = a.cols, 320, a.cols * a.rows
    sheets = 0
    for i in range(0, len(files), per):
        chunk = files[i:i + per]
        thumbs = []
        for f in chunk:
            im = Image.open(f).convert("RGB")
            im.thumbnail((cell_w, cell_w))
            thumbs.append((f.name, im))
        cell_h = max(im.height for _, im in thumbs) + 22
        rows = (len(thumbs) + cols - 1) // cols
        sheet = Image.new("RGB", (cols * cell_w, rows * cell_h), "white")
        dr = ImageDraw.Draw(sheet)
        for k, (name, im) in enumerate(thumbs):
            x, y = (k % cols) * cell_w, (k // cols) * cell_h
            sheet.paste(im, (x, y))
            dr.text((x + 3, y + im.height + 4), name[:48], fill="black")
        sheets += 1
        sheet.save(fdir / f"contact_{sheets:02d}.jpg", quality=85)
    print(f"{sheets} contact sheet(s) em {fdir}/contact_NN.jpg ({len(files)} frames)")
    return 0


# ---------- fase 7: docx ----------
def cmd_docx(a):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    d = load_doc(a)
    w = work(a.name)
    fdir = Path(d["frames_dir"])
    NAVY, GRAY = RGBColor(0x1F, 0x3A, 0x93), RGBColor(0x88, 0x88, 0x88)
    img_w = Inches(d.get("img_width_in", 2.35))

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    for sname, size, color in (("Title", 26, "17365D"), ("Heading 1", 14, "366091"),
                               ("Heading 2", 13, "4F81BD")):
        st = doc.styles[sname]
        st.font.name, st.font.size = "Calibri", Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        if sname != "Title":
            st.font.bold = True
    doc.styles["Normal"].font.name = "Calibri"

    def small(text, italic=True, size=9.5):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.italic, r.font.size = italic, Pt(size)
        return p

    def fala(f, lang):
        p = doc.add_paragraph()
        r = p.add_run(f"{f['falante']}: ")
        r.bold, r.font.color.rgb = True, NAVY
        p.add_run(f[lang])

    def bloco_h1(b):
        doc.add_heading(f"BLOCO {b['n']} — {b['titulo']}  [{b['inicio']} – {b['fim']}]", level=1)

    def mini_h2(b, k, mb):
        doc.add_heading(f"{b['n']}.{k}   ⏱ {mb['inicio']} → {mb['fim']}", level=2)

    # cabeçalho
    doc.add_heading(d["titulo"], level=0)
    small(d["fonte"])
    small(d["nota_prints"])
    small("Falantes: " + " · ".join(d["falantes"]), italic=False)

    doc.add_heading(f"Estrutura da VSL ({len(d['blocos'])} blocos)", level=1)
    for b in d["blocos"]:
        small(f"BLOCO {b['n']} — {b['titulo']}  —  {b['inicio']} → {b['fim']}", italic=False, size=10)

    # parte 1: transcrição original
    doc.add_heading(d.get("titulo_parte1", "PARTE 1 — TRANSCRIÇÃO EM INGLÊS POR BLOCOS"), level=0)
    for b in d["blocos"]:
        bloco_h1(b)
        for k, mb in enumerate(b["mini"], 1):
            mini_h2(b, k, mb)
            for f in mb["falas"]:
                fala(f, "en")
    doc.add_paragraph()

    # parte 2: tradução + prints
    doc.add_heading("LER AQUI", level=0)
    doc.add_heading(d.get("titulo_parte2", "PARTE 2 — TRADUÇÃO PT-BR + PRINTS"), level=0)
    n_frames, n_caps = 0, 0
    for b in d["blocos"]:
        bloco_h1(b)
        for k, mb in enumerate(b["mini"], 1):
            mini_h2(b, k, mb)
            prints = mb.get("prints", [])
            if prints:
                rows = (len(prints) + 1) // 2
                t = doc.add_table(rows=rows, cols=2)
                for i, pr in enumerate(prints):
                    cell = t.cell(i // 2, i % 2)
                    pimg = cell.paragraphs[0]
                    pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fpath = fdir / pr["arquivo"]
                    if fpath.is_file():
                        pimg.add_run().add_picture(str(fpath), width=img_w)
                        n_frames += 1
                    else:
                        pimg.add_run(f"[FRAME INDISPONÍVEL em {pr['ts']}]").italic = True
                    pc = cell.add_paragraph()
                    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    rc = pc.add_run(f"📷 {pr['ts']} — {pr['legenda']}")
                    rc.italic, rc.font.size = True, Pt(8.5)
                    rc.add_break()
                    rf = pc.add_run(pr["arquivo"])
                    rf.font.size, rf.font.color.rgb = Pt(7), GRAY
                    n_caps += 1
                doc.add_paragraph()
            for f in mb["falas"]:
                fala(f, "pt")

    # Notas de adaptação (Fase 5 do PROMPT): sempre presentes, salvo tradução fiel
    notas = d.get("notas")
    if notas is not None:
        doc.add_page_break()
        doc.add_heading("NOTAS DE ADAPTAÇÃO", level=1)
        for titulo, chave in (("Adaptações culturais", "adaptacoes_culturais"),
                              ("Personagens renomeados", "personagens_renomeados"),
                              ("Conversões de moeda/medida", "conversoes"),
                              ("Pontos com dúvida [?? ??]", "duvidas"),
                              ("Erros de transcrição corrigidos", "correcoes_asr"),
                              ("Inconsistências no original", "inconsistencias")):
            doc.add_heading(titulo, level=2)
            itens = notas.get(chave) or []
            if not itens:
                doc.add_paragraph("Nenhuma.")
            for it in itens:
                doc.add_paragraph(it, style="List Bullet")

    out = w / d.get("arquivo_saida", f"{a.name}_VSL__Transcricao_EN__Traducao_PTBR__Prints.docx")
    doc.save(out)
    n_mini = sum(len(b["mini"]) for b in d["blocos"])
    print(f"Doc: {out} — {len(d['blocos'])} blocos, {n_mini} mini-blocos, "
          f"{n_frames} frames embutidos / {n_caps} legendas")
    return 0


# ---------- fase 8: verificação ----------
def cmd_verify(a):
    d = load_doc(a)
    fdir = Path(d["frames_dir"])
    errs = []
    last_end, n_prints, n_falas = -1, 0, 0
    for b in d["blocos"]:
        if ts_to_sec(b["inicio"]) < last_end - 1:
            errs.append(f"BLOCO {b['n']} começa ({b['inicio']}) antes do fim do bloco anterior")
        prev_mini_end = ts_to_sec(b["inicio"]) - 1
        for k, mb in enumerate(b["mini"], 1):
            s, e = ts_to_sec(mb["inicio"]), ts_to_sec(mb["fim"])
            if s < prev_mini_end - 1:
                errs.append(f"{b['n']}.{k} começa antes do mini-bloco anterior")
            if e < s:
                errs.append(f"{b['n']}.{k} termina antes de começar")
            prev_mini_end = e
            if not mb.get("falas"):
                errs.append(f"{b['n']}.{k} sem falas")
            for f in mb.get("falas", []):
                n_falas += 1
                for lang in ("en", "pt"):
                    if not f.get(lang, "").strip():
                        errs.append(f"{b['n']}.{k} fala de {f.get('falante')} sem texto '{lang}'")
                if f.get("falante") not in d["falantes"]:
                    errs.append(f"{b['n']}.{k} falante '{f.get('falante')}' fora da lista de falantes")
            for pr in mb.get("prints", []):
                n_prints += 1
                t = ts_to_sec(pr["ts"])
                if not (s - 2 <= t <= e + 2):
                    errs.append(f"{b['n']}.{k} print {pr['ts']} fora do intervalo {mb['inicio']}→{mb['fim']}")
                if not (fdir / pr["arquivo"]).is_file():
                    errs.append(f"{b['n']}.{k} arquivo ausente: {pr['arquivo']}")
                if not re.match(r"^b\d{2}_[a-z0-9_]+_\d{2}-\d{2}-\d{2}\.jpg$", pr["arquivo"]):
                    errs.append(f"{b['n']}.{k} nome fora do padrão bNN_descricao_HH-MM-SS.jpg: {pr['arquivo']}")
                if not pr.get("legenda", "").strip():
                    errs.append(f"{b['n']}.{k} print {pr['ts']} sem legenda")
        last_end = ts_to_sec(b["fim"])
    key = os.environ.get("ASSEMBLYAI_API_KEY", "")
    if key and key in json.dumps(d):
        errs.append("A CHAVE DA API APARECE NO doc.json")
    print(f"{len(d['blocos'])} blocos, {n_falas} falas, {n_prints} prints com legenda.")
    if errs:
        print(f"{len(errs)} problema(s):")
        for e in errs:
            print(" -", e)
        return 1
    print("Verificação OK.")
    return 0


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("check").set_defaults(fn=cmd_check)
    for name, fn in (("prep", cmd_prep), ("audio", cmd_audio)):
        s = sub.add_parser(name); s.add_argument("video"); s.add_argument("--name", required=True)
        s.set_defaults(fn=fn)
    s = sub.add_parser("transcribe"); s.add_argument("--name", required=True)
    s.add_argument("--lang", default="en"); s.set_defaults(fn=cmd_transcribe)
    s = sub.add_parser("frames"); s.add_argument("video"); s.add_argument("--name", required=True)
    s.add_argument("--doc"); s.add_argument("--every", type=int, default=60)
    s.add_argument("--extra", nargs="*", default=[]); s.set_defaults(fn=cmd_frames)
    s = sub.add_parser("sheet"); s.add_argument("--name", required=True)
    s.add_argument("--cols", type=int, default=4); s.add_argument("--rows", type=int, default=5)
    s.set_defaults(fn=cmd_sheet)
    for name, fn in (("docx", cmd_docx), ("verify", cmd_verify)):
        s = sub.add_parser(name); s.add_argument("--name", required=True); s.add_argument("--doc")
        s.set_defaults(fn=fn)
    a = p.parse_args()
    sys.exit(a.fn(a))


if __name__ == "__main__":
    main()
